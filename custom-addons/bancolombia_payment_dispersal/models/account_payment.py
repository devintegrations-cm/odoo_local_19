# -*- coding: utf-8 -*-
# Part of Odoo. See LICENSE file for full copyright and licensing details.
import logging
import base64
import io
import unidecode
from openpyxl import Workbook
from datetime import date, datetime, timedelta, time

from odoo import api, fields, models, _, Command
from odoo.tools.safe_eval import safe_eval
from odoo.tools import DEFAULT_SERVER_DATETIME_FORMAT


from odoo.addons.account_payment_dispersion.tools import _column_name_field

_logger = logging.getLogger(__name__)


class AccountPayment(models.Model):

    _inherit = 'account.payment'

    @staticmethod
    def _get_bancolombia_type_of_transaction():
        return [
            ('23', _('Pre-notification Checking Account')),
            ('25', _('Cash Payment')),
            ('27', _('Deposit to Checking Account')),
            ('33', _('Pre-notify savings account')),
            ('36', _('Manager\'s Check Payment')),
            ('37', _('Deposit to Savings Account')),
            ('40', _('Secure Cash (Visa Payments)')),
            ('52', _('Electronic deposit payment')),
            ('53', _('Pre-notification of Electronic Deposit'))
        ]

    # Mismo comportamiento que en 17: sin rama que vacíe el campo, para no
    # pisar los tipos cargados a mano (25/36/40 no llevan cuenta destino).
    @api.depends('partner_id', 'partner_bank_id')
    def _compute_bancolombia_type_of_transaction(self):
        for pay in self:
            if pay.payment_type == 'outbound' and pay.type == 'bank' and pay.partner_id and pay.partner_bank_id and pay.partner_bank_id.bancolombia_type_of_transaction:
                pay.bancolombia_type_of_transaction = pay.partner_bank_id.bancolombia_type_of_transaction

    bancolombia_type_of_transaction = fields.Selection(
        selection='_get_bancolombia_type_of_transaction', 
        string='Type of transaction',
        readonly=False, store=True, tracking=True,
        compute='_compute_bancolombia_type_of_transaction',)

    def action_generate_bancolombia_dispersion_file(self, type_of_payment='220'):


        if len(list(set(self.mapped('journal_id')))) > 1:
            msg = _("The payment dispersions must be for a single journal and in this case it has the following journals: %s", ', '.join(self.mapped('journal_id').mapped('name')))
            return self.show_meesage(msg)

        if self.filtered(lambda payment: not payment.journal_id.bank_account_id):
            msg = _("The journal %s does not have a source account set up", self.journal_id.name)
            return self.show_meesage(msg)

        if self.filtered(lambda payment: not payment.bancolombia_type_of_transaction):
            msg = _("No transaction type configured: %s", self.filtered(lambda payment: not payment.bancolombia_type_of_transaction)._bancolombia_payment_names())
            return self.show_meesage(msg)
        
        cxt = {
            'env': self.env,
            'date': date,
            'datetime': datetime,
            'timedelta': timedelta,
            'time': time,

        }

        wb = Workbook()
        ws = wb.active

        bancolombia_dispersion_fields = self.env['bancolombia.payment_dispersal_field'].search(
            [], order="sequence asc")
        data = list()

        for rec in self:

            if rec.bancolombia_type_of_transaction not in ('25', '36', '40') and rec.filtered(lambda payment: not payment.partner_bank_id):
                msg = _("One or more of the payments has no target account: %s", self.filtered(lambda payment: not payment.partner_bank_id)._bancolombia_payment_names())
                return self.show_meesage(msg)

            if rec.bancolombia_type_of_transaction not in ('25', '36', '40') and rec.filtered(lambda payment: (not payment.partner_bank_id.bank_id or not payment.partner_bank_id.bank_id.code)):
                msg = _("No bank or bank code configured: %s", self.filtered(lambda payment: (not payment.partner_bank_id.bank_id or not payment.partner_bank_id.bank_id.code))._bancolombia_payment_names())
                return self.show_meesage(msg)


            if bancolombia_dispersion_fields:

                names = bancolombia_dispersion_fields.mapped('name')

                cxt.update({'object': rec})

                ws = _column_name_field(names, ws)
                map_header_vals = {
                    'A':["NIT PAGADOR", rec.company_id.vat[0:-2] if rec.company_id.vat else ""],
                    'B':["TIPO DE PAGO",type_of_payment],
                    'C':["APLICACION", "I"],
                    'D':["SECUENCA DE ENVÍO"],
                    'E':["NRO CUENTA A DEBITAR", rec.journal_id.bank_account_id.acc_number if rec.journal_id.bank_account_id else ""],
                    'F':["TIPO DE CUENTA A DEBITAR", rec._get_account_type(rec.journal_id.bank_account_id.type_of_account) if rec.journal_id.bank_account_id else ""],
                    'G':["DESCRIPCIÓN DEL PAGO"]
                }
                for col, content in map_header_vals.items():
                    for ix, info in enumerate(content):
                        pass#ws[f'{col}{str(ix+1)}']=info

                ws['A1'] = 'NIT PAGADOR'
                if rec.company_id.vat:
                    ws['A2'] = rec.company_id.vat[0:-2]

                ws['B1'] = 'TIPO DE PAGO'
                ws['B2'] = type_of_payment

                ws['C1'] = 'APLICACIÓN'
                ws['C2'] = 'I'

                ws['D1'] = 'SECUENCIA DE ENVIÓ'
                ws['D2'] = 'A1'

                ws['E1'] = 'NRO CUENTA A DEBITAR'
                if rec.journal_id.bank_account_id:
                    ws['E2'] = rec.journal_id.bank_account_id.acc_number

                ws['F1'] = 'TIPO DE CUENTA A DEBITAR'
                if rec.journal_id.bank_account_id:
                    ws['F2'] = rec._get_account_type(rec.journal_id.bank_account_id.type_of_account)

                ws['G1'] = 'DESCRIPCIÓN DEL PAGO'
                
                rows = list()
                for bancolombia_dispersion_field in bancolombia_dispersion_fields:
                    safe_eval(bancolombia_dispersion_field.value,
                              cxt, mode="exec")
                    result = cxt.get('result')

                    if result:
                        rows.append(result)
                data.append(rows)

        for row in data:
            ws.append(row)

        output = io.BytesIO()
        wb.save(output)

        file_name = f'bancolombia_dispersion_{type_of_payment}_{datetime.now().strftime(DEFAULT_SERVER_DATETIME_FORMAT)}'
        dispersion_payment = self._create_bancolombia_payment_dispersal(file_name, output)

        return {
            'name': file_name,
            'type': 'ir.actions.act_url',
            'url': "web/content/?model=account.payment_dispersal&id=" + str(dispersion_payment.id) + "&filename_field=binary_file_name&field=binary_file&download=true&filename=" + f'{file_name}.xlsx',
            'target': 'self',
        }

    def _create_bancolombia_payment_dispersal(self, file_name, output):
        dispersion_payment = self.env['account.payment_dispersal'].create(
            {'binary_file': base64.b64encode(output.getvalue()), 'binary_file_name': f'{file_name}.xlsx', 'name': file_name, 'account_payment_ids': [Command.set(self.ids)], 'journal_id': self.mapped('journal_id')[0].id})
        return dispersion_payment


    def _bancolombia_payment_names(self):
        # En 19 un pago en borrador no tiene número (name vacío; en 17 era '/').
        return ', '.join(
            pay.name or f"{pay.display_name} ({pay.partner_id.name or pay.id})"
            for pay in self
        )

    def _get_account_type(self, account_type):
        if account_type == 'current_account':
            return 'D'
        elif account_type == 'savings_account':
            return 'S'

    def show_meesage(self, msg):
        return {
                'type': 'ir.actions.client',
                'tag': 'display_notification',
                'params': {
                    'type': 'info',
                    'sticky': True,
                    'message': _(msg),
                }
            }